"""Report Export API endpoints.

Provides PDF and Excel report generation for:
- Individual candidate reports
- Assessment overview reports
- Data export for analytics
"""

import io
import json
from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import select, func
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.permissions import require_roles
from app.db.session import get_db
from app.models.assessment import Assessment, AssessmentQuestion
from app.models.candidate import CandidateSession
from app.models.question import Question
from app.models.response import CandidateResponse
from app.models.user import User

router = APIRouter(tags=["reports"])


# ---------------------------------------------------------------------------
# GET /reports/candidate/{session_id}/pdf
# ---------------------------------------------------------------------------

@router.get("/candidate/{session_id}/pdf")
async def export_candidate_pdf(
    session_id: UUID,
    user: User = Depends(require_roles("admin", "hr", "tech")),
    db: AsyncSession = Depends(get_db),
):
    """Generate a PDF report for a candidate session."""
    session = await db.get(CandidateSession, session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")

    assessment = await db.get(Assessment, session.assessment_id)

    # Load responses
    result = await db.execute(
        select(CandidateResponse)
        .where(CandidateResponse.session_id == session_id)
        .options(selectinload(CandidateResponse.question))
    )
    responses = result.scalars().all()

    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.units import inch
    except ImportError:
        raise HTTPException(
            status_code=503,
            detail="PDF generation requires 'reportlab'. Install it with: pip install reportlab"
        )

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=0.5*inch, bottomMargin=0.5*inch)
    styles = getSampleStyleSheet()
    story = []

    # Title
    title_style = ParagraphStyle("Title", parent=styles["Heading1"], fontSize=18, spaceAfter=12)
    story.append(Paragraph("Candidate Assessment Report", title_style))
    story.append(Spacer(1, 12))

    # Candidate info
    info_data = [
        ["Candidate:", session.candidate_name],
        ["Email:", session.candidate_email],
        ["Assessment:", assessment.title if assessment else "N/A"],
        ["Status:", session.status],
        ["Score:", f"{session.score_pct:.1f}%" if session.score_pct else "N/A"],
        ["Result:", "PASSED" if session.is_passed else "NOT PASSED" if session.is_passed is not None else "PENDING"],
    ]
    info_table = Table(info_data, colWidths=[1.5*inch, 4.5*inch])
    info_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(info_table)
    story.append(Spacer(1, 20))

    # Question results table
    story.append(Paragraph("Question Results", styles["Heading2"]))
    story.append(Spacer(1, 8))

    q_headers = ["#", "Question", "Type", "Score", "Max"]
    q_data = [q_headers]
    for i, resp in enumerate(responses, 1):
        q = resp.question
        score = resp.final_score or resp.auto_score or "-"
        q_data.append([
            str(i),
            (q.title[:40] + "..." if len(q.title) > 40 else q.title) if q else "N/A",
            (q.type.value if hasattr(q.type, "value") else q.type) if q else "-",
            str(score),
            str(resp.max_score or q.max_score if q else "-"),
        ])

    q_table = Table(q_data, colWidths=[0.4*inch, 3*inch, 1*inch, 0.8*inch, 0.8*inch])
    q_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563eb")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f0f4ff")]),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(q_table)

    doc.build(story)
    buffer.seek(0)

    filename = f"candidate_report_{session.candidate_name.replace(' ', '_')}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


# ---------------------------------------------------------------------------
# GET /reports/assessment/{assessment_id}/pdf
# ---------------------------------------------------------------------------

@router.get("/assessment/{assessment_id}/pdf")
async def export_assessment_pdf(
    assessment_id: UUID,
    user: User = Depends(require_roles("admin", "hr", "tech")),
    db: AsyncSession = Depends(get_db),
):
    """Generate a PDF overview report for an assessment."""
    assessment = await db.get(Assessment, assessment_id)
    if not assessment:
        raise HTTPException(status_code=404, detail="Assessment not found")

    # Load sessions
    result = await db.execute(
        select(CandidateSession)
        .where(CandidateSession.assessment_id == assessment_id)
        .order_by(CandidateSession.score_pct.desc().nulls_last())
    )
    sessions = result.scalars().all()

    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.units import inch
    except ImportError:
        raise HTTPException(status_code=503, detail="Install reportlab: pip install reportlab")

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph(f"Assessment Report: {assessment.title}", styles["Heading1"]))
    story.append(Spacer(1, 12))

    # Stats
    completed = [s for s in sessions if s.status == "completed"]
    avg_score = sum(s.score_pct or 0 for s in completed) / len(completed) if completed else 0
    pass_rate = sum(1 for s in completed if s.is_passed) / len(completed) * 100 if completed else 0

    stats_data = [
        ["Total Candidates:", str(len(sessions))],
        ["Completed:", str(len(completed))],
        ["Average Score:", f"{avg_score:.1f}%"],
        ["Pass Rate:", f"{pass_rate:.1f}%"],
    ]
    stats_table = Table(stats_data, colWidths=[2*inch, 4*inch])
    stats_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
    ]))
    story.append(stats_table)
    story.append(Spacer(1, 20))

    # Rankings table
    story.append(Paragraph("Candidate Rankings", styles["Heading2"]))
    headers = ["Rank", "Candidate", "Score", "Status", "Result"]
    data = [headers]
    for i, sess in enumerate(sessions, 1):
        data.append([
            str(i),
            sess.candidate_name,
            f"{sess.score_pct:.1f}%" if sess.score_pct else "-",
            sess.status,
            "PASS" if sess.is_passed else "FAIL" if sess.is_passed is not None else "-",
        ])

    t = Table(data, colWidths=[0.5*inch, 2.5*inch, 1*inch, 1*inch, 1*inch])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563eb")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f0f4ff")]),
    ]))
    story.append(t)
    doc.build(story)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=assessment_{assessment_id}.pdf"},
    )


# ---------------------------------------------------------------------------
# GET /reports/assessment/{assessment_id}/excel
# ---------------------------------------------------------------------------

@router.get("/assessment/{assessment_id}/excel")
async def export_assessment_excel(
    assessment_id: UUID,
    user: User = Depends(require_roles("admin", "hr", "tech")),
    db: AsyncSession = Depends(get_db),
):
    """Export assessment data as Excel spreadsheet."""
    assessment = await db.get(Assessment, assessment_id)
    if not assessment:
        raise HTTPException(status_code=404, detail="Assessment not found")

    result = await db.execute(
        select(CandidateSession)
        .where(CandidateSession.assessment_id == assessment_id)
        .order_by(CandidateSession.score_pct.desc().nulls_last())
    )
    sessions = result.scalars().all()

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(status_code=503, detail="Install openpyxl: pip install openpyxl")

    wb = Workbook()
    ws = wb.active
    ws.title = "Candidate Results"

    # Header styling
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")

    headers = ["Rank", "Candidate Name", "Email", "Status", "Score %", "Total Score",
               "Max Score", "Pass/Fail", "Violations", "Started", "Completed"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row, sess in enumerate(sessions, 2):
        ws.cell(row=row, column=1, value=row - 1)
        ws.cell(row=row, column=2, value=sess.candidate_name)
        ws.cell(row=row, column=3, value=sess.candidate_email)
        ws.cell(row=row, column=4, value=sess.status)
        ws.cell(row=row, column=5, value=round(sess.score_pct, 1) if sess.score_pct else None)
        ws.cell(row=row, column=6, value=round(sess.total_score, 2) if sess.total_score else None)
        ws.cell(row=row, column=7, value=round(sess.total_max_score, 2) if sess.total_max_score else None)
        ws.cell(row=row, column=8, value="PASS" if sess.is_passed else "FAIL" if sess.is_passed is not None else "")
        ws.cell(row=row, column=9, value=sess.proctoring_violations or 0)
        ws.cell(row=row, column=10, value=sess.started_at.isoformat() if sess.started_at else "")
        ws.cell(row=row, column=11, value=sess.completed_at.isoformat() if sess.completed_at else "")

    # Auto-width
    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 30)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=assessment_{assessment_id}.xlsx"},
    )
